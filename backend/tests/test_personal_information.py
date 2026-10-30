from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import get_settings


def test_profile_patch_preserves_unmentioned_fields(client: TestClient) -> None:
    first = client.patch(
        "/api/profile",
        json={"background": "Studying software development", "interests": ["design", "coding"]},
    )
    second = client.patch("/api/profile", json={"motivations": "Build useful tools"})

    assert first.status_code == 200
    assert second.status_code == 200
    profile = second.json()
    assert profile["background"] == "Studying software development"
    assert profile["interests"] == ["design", "coding"]
    assert profile["motivations"] == "Build useful tools"


def test_career_and_academic_goals_can_be_created_updated_and_deleted(
    client: TestClient,
) -> None:
    career = client.post(
        "/api/goals/career",
        json={
            "title": "Explore backend engineering",
            "required_skills": ["Python", "SQL"],
            "importance": 4,
        },
    )
    academic = client.post(
        "/api/goals/academic",
        json={"title": "Finish databases course", "course": "Databases", "target_date": "2027-05-01"},
    )

    assert career.status_code == 201
    assert career.json()["required_skills"] == ["Python", "SQL"]
    assert academic.status_code == 201
    assert academic.json()["target_date"] == "2027-05-01"

    updated = client.patch(
        f"/api/goals/career/{career.json()['id']}",
        json={"status": "paused"},
    )
    assert updated.status_code == 200
    assert updated.json()["status"] == "paused"

    deleted = client.delete(f"/api/goals/academic/{academic.json()['id']}")
    assert deleted.status_code == 204
    assert client.get("/api/goals/academic").json() == []


def test_principles_support_editing_and_status_changes(client: TestClient) -> None:
    created = client.post(
        "/api/principles",
        json={
            "title": "Consistency over perfection",
            "category": "learning",
            "contexts": ["study", "projects"],
        },
    )
    assert created.status_code == 201
    principle_id = created.json()["id"]

    updated = client.patch(
        f"/api/principles/{principle_id}",
        json={"active": False, "reminder": "Take the next small step"},
    )
    assert updated.status_code == 200
    assert updated.json()["active"] is False
    assert updated.json()["reminder"] == "Take the next small step"
    assert client.get("/api/principles").json()[0]["title"] == "Consistency over perfection"


def _sample_resume() -> bytes:
    document = Document()
    document.add_paragraph("Alex Example")
    document.add_paragraph("alex@example.test")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Junior developer at Example Studio, 2019 - 2021")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python and SQL")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_resume_upload_extracts_unverified_facts_and_preserves_original(
    client: TestClient,
    tmp_path: Path,
) -> None:
    resume_bytes = _sample_resume()
    upload = client.post(
        "/api/resume/upload",
        files={
            "file": (
                "../../alex-cv.docx",
                resume_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert upload.status_code == 201
    document = upload.json()
    assert document["original_filename"] == "alex-cv.docx"
    assert document["version"] == 1
    assert document["extraction_status"] == "pending_review"
    assert document["facts"]
    assert all(fact["verified"] is False for fact in document["facts"])
    assert any(fact["category"] == "experience" for fact in document["facts"])
    assert any(fact["category"] == "skill" for fact in document["facts"])
    assert not any(
        fact["category"] == "contact" and fact["value"].startswith("Phone:")
        for fact in document["facts"]
    )

    download = client.get(f"/api/resume/{document['id']}/download")
    assert download.status_code == 200
    assert download.content == resume_bytes
    assert download.headers["content-disposition"].startswith("attachment;")

    fact = document["facts"][0]
    verified = client.patch(
        f"/api/resume/{document['id']}/facts/{fact['id']}",
        json={"value": "Reviewed by the user", "verified": True},
    )
    assert verified.status_code == 200
    assert verified.json()["value"] == "Reviewed by the user"
    assert verified.json()["verified"] is True
    assert client.get(f"/api/resume/{document['id']}").json()["extraction_status"] == "pending_review"
    for remaining_fact in document["facts"][1:]:
        response = client.patch(
            f"/api/resume/{document['id']}/facts/{remaining_fact['id']}",
            json={"verified": True},
        )
        assert response.status_code == 200
    assert client.get(f"/api/resume/{document['id']}").json()["extraction_status"] == "reviewed"
    assert len(list((tmp_path / "files").iterdir())) == 1


def test_resume_upload_rejects_mismatched_file_content(client: TestClient) -> None:
    response = client.post(
        "/api/resume/upload",
        files={"file": ("resume.pdf", b"not a PDF", "application/pdf")},
    )
    assert response.status_code == 422
    assert response.json()["detail"] == "The uploaded file is not a valid PDF."


def test_resume_upload_enforces_configured_file_size_limit(
    client: TestClient,
    monkeypatch,
) -> None:
    monkeypatch.setattr(get_settings(), "max_resume_size_mb", 1)
    response = client.post(
        "/api/resume/upload",
        files={"file": ("resume.pdf", b"x" * (1024 * 1024 + 1), "application/pdf")},
    )
    assert response.status_code == 413


def test_resume_deletion_removes_original_and_extracted_details(
    client: TestClient,
    tmp_path: Path,
) -> None:
    uploaded = client.post(
        "/api/resume/upload",
        files={
            "file": (
                "resume.docx",
                _sample_resume(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document = uploaded.json()
    storage_dir = tmp_path / "files"
    stored_file = next(storage_dir.iterdir())
    deleted = client.delete(f"/api/resume/{document['id']}")

    assert deleted.status_code == 204
    assert not stored_file.exists()
    assert client.get("/api/resume").json() == []
